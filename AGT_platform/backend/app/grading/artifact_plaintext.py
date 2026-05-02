"""
Convert graded **artifact bytes** (submission or blank template) into plain text for LLM/RAG.

Used by :mod:`app.grading.submission_text` and :mod:`app.grading.multimodal.llm_triplet_three_source`.
Optional formats (``docx``, ``xlsx``) use soft imports so the backend still starts if an extra
library is not installed; those cases return a short placeholder string.
"""

from __future__ import annotations

import csv
import io
import logging
from typing import Final

from app.grading.tools import extract_from_ipynb, extract_text_from_pdf, transcribe_video_stub

_log = logging.getLogger(__name__)

# Stable key order for concatenated submission views.
ARTIFACT_KEY_ORDER: Final[tuple[str, ...]] = (
    "ipynb",
    "pdf",
    "docx",
    "txt",
    "md",
    "py",
    "csv",
    "xlsx",
    "mp4",
    "mp3",
    "wav",
    "m4a",
    "webm",
)

def bytes_with_suffix_to_plain(data: bytes, suffix: str) -> str:
    """
    Decode ``data`` using the file extension ``suffix`` (e.g. ``".pdf"``, ``".py"``).

    For unknown suffixes, falls back to UTF-8 decode with replacement.
    """
    if not data or not data.strip():
        return ""
    suf = (suffix or "").lower().strip()
    if not suf.startswith("."):
        suf = "." + suf if suf else ""
    try:
        if suf == ".ipynb":
            nb = extract_from_ipynb(data)
            code = (nb.get("code") or "").strip()
            md = (nb.get("markdown") or "").strip()
            parts: list[str] = []
            if md:
                parts.append(f"=== MARKDOWN ===\n{md}")
            if code:
                parts.append(f"=== CODE CELLS ===\n{code}")
            return "\n\n".join(parts).strip()
        if suf == ".pdf":
            return extract_text_from_pdf(data).strip()
        if suf == ".docx":
            return _docx_bytes_to_plain(data)
        if suf in (".txt", ".md", ".py"):
            return data.decode("utf-8", errors="replace").strip()
        if suf == ".csv":
            return _csv_bytes_to_plain(data)
        if suf == ".xlsx":
            return _xlsx_bytes_to_plain(data)
        if suf in (".mp4", ".mp3", ".wav", ".m4a", ".webm"):
            tr = transcribe_video_stub(data)
            return f"=== VIDEO / AUDIO ({len(data)} bytes) ===\n{tr}".strip()
        return data.decode("utf-8", errors="replace").strip()
    except Exception:
        _log.debug("bytes_with_suffix_to_plain failed suffix=%s", suf, exc_info=True)
        return ""


def _docx_bytes_to_plain(data: bytes) -> str:
    try:
        import docx  # type: ignore[import-untyped]
    except ImportError:
        return "[.docx present; install python-docx for text extraction]"
    try:
        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs if (p.text or "").strip()]
        return "\n".join(lines).strip()
    except Exception:
        _log.warning("docx plaintext extraction failed", exc_info=True)
        return "[.docx binary; text extraction failed]"


def _xlsx_bytes_to_plain(data: bytes) -> str:
    try:
        import openpyxl  # type: ignore[import-untyped]

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(
                    "\t".join("" if c is None else str(c) for c in row)
                )
            parts.append("\n".join(rows))
        return "\n\n".join(parts).strip()
    except ImportError:
        pass
    except Exception:
        _log.debug("openpyxl xlsx read failed", exc_info=True)
    try:
        import pandas as pd  # type: ignore[import-untyped]

        frames: list[str] = []
        xl = pd.ExcelFile(io.BytesIO(data))
        for name in xl.sheet_names:
            df = pd.read_excel(io.BytesIO(data), sheet_name=name, header=None)
            frames.append(f"=== SHEET {name} ===\n{df.to_csv(index=False, sep='\t')}")
        return "\n\n".join(frames).strip()
    except Exception:
        _log.debug("pandas xlsx read failed", exc_info=True)
    return "[.xlsx present; install openpyxl or pandas for text extraction]"


def _csv_bytes_to_plain(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    sio = io.StringIO(text)
    try:
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        return text.strip()
    sio.seek(0)
    rows_out: list[str] = []
    reader = csv.reader(sio, dialect)
    for row in reader:
        rows_out.append("\t".join(row))
    return "\n".join(rows_out).strip()


def single_artifact_key_to_plain(key: str, raw: bytes) -> str:
    """Map ingestion artifact key (e.g. ``pdf``) to extracted text."""
    if not raw:
        return ""
    k = (key or "").strip().lower()
    try:
        if k == "ipynb":
            nb = extract_from_ipynb(raw)
            code = (nb.get("code") or "").strip()
            md = (nb.get("markdown") or "").strip()
            parts: list[str] = []
            if md:
                parts.append(f"=== NOTEBOOK MARKDOWN ({k}) ===\n{md}")
            if code:
                parts.append(f"=== NOTEBOOK CODE ({k}) ===\n{code}")
            return "\n\n".join(parts).strip()
        if k == "pdf":
            txt = extract_text_from_pdf(raw).strip()
            return f"=== PDF TEXT ===\n{txt}" if txt else ""
        if k == "docx":
            body = _docx_bytes_to_plain(raw)
            return f"=== DOCX ===\n{body}" if body else ""
        if k in ("txt", "md"):
            return (
                f"=== {k.upper()} ===\n"
                + raw.decode("utf-8", errors="replace").strip()
            )
        if k == "py":
            return (
                "=== PYTHON SOURCE ===\n"
                + raw.decode("utf-8", errors="replace").strip()
            )
        if k == "csv":
            body = _csv_bytes_to_plain(raw)
            return f"=== CSV ===\n{body}" if body else ""
        if k == "xlsx":
            body = _xlsx_bytes_to_plain(raw)
            return f"=== XLSX ===\n{body}" if body else ""
        if k in ("mp4", "mp3", "wav", "m4a", "webm"):
            tr = transcribe_video_stub(raw)
            return f"=== VIDEO / AUDIO ({len(raw)} bytes) ===\n{tr}"
    except Exception:
        _log.warning("single_artifact_key_to_plain failed key=%s", k, exc_info=True)
        return f"=== ERROR parsing artifact {k} ===\n"
    return ""


def artifacts_to_concatenated_plain(artifacts: dict[str, bytes]) -> str:
    """Concatenate normalized text from artifact dict keys in :data:`ARTIFACT_KEY_ORDER`."""
    if not isinstance(artifacts, dict) or not artifacts:
        return ""
    chunks: list[str] = []
    for key in ARTIFACT_KEY_ORDER:
        raw = artifacts.get(key)
        if not isinstance(raw, (bytes, bytearray)) or not raw:
            continue
        block = single_artifact_key_to_plain(key, bytes(raw))
        if block.strip():
            chunks.append(block)
    if not chunks:
        return ""
    return "\n\n".join(chunks)


def infer_modality_from_artifact_keys(artifacts: dict[str, bytes]) -> str:
    """Return a :class:`~app.grading.multimodal.schemas.Modality` value string for hints."""
    if not isinstance(artifacts, dict) or not artifacts:
        return "unknown"
    keys = {k.lower() for k, v in artifacts.items() if isinstance(v, (bytes, bytearray)) and v}
    if "ipynb" in keys:
        return "notebook"
    if keys & {"mp4", "mp3", "wav", "m4a", "webm"}:
        return "video_oral"
    if "pdf" in keys or "docx" in keys or "txt" in keys or "md" in keys:
        return "written"
    if "py" in keys:
        return "code"
    if "csv" in keys or "xlsx" in keys:
        return "programming_analysis"
    if len(keys) > 1:
        return "mixed"
    return "unknown"
